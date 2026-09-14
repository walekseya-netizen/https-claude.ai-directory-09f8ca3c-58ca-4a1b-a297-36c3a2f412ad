"""Тесты выгрузки форм в XLSX."""

from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import load_workbook

from app.domain.calc import build_act, build_certificate_from_acts
from app.domain.models import KS2Input, KS3FromActs
from app.export.xlsx_ks2 import render_ks2
from app.export.xlsx_ks3 import render_ks3
from tests.conftest import ks2_payload


def load(content: bytes):
    return load_workbook(BytesIO(content)).active


def cell_values(sheet) -> list:
    return [cell.value for row in sheet.iter_rows() for cell in row if cell.value not in (None, "")]


def test_ks2_sheet_structure():
    act = build_act(KS2Input.model_validate(ks2_payload()))

    sheet = load(render_ks2(act))
    values = cell_values(sheet)

    assert sheet.title == "КС-2"
    assert "АКТ О ПРИЕМКЕ ВЫПОЛНЕННЫХ РАБОТ" in values
    assert "Унифицированная форма № КС-2" in values
    assert "0322005" in values
    assert "ООО «Подрядчик», г. Москва, ул. Строителей, д. 5" in values
    assert "Жилой комплекс «Северный», г. Москва, ул. Полярная, д. 10" in values
    assert "Итого" in values
    assert "Сумма НДС (20%)" in values
    assert "Всего с учетом НДС" in values
    assert float(act.totals.gross) in values
    assert "Устройство монолитных железобетонных стен" in values


def test_ks2_line_amounts_and_coefficient_note():
    act = build_act(KS2Input.model_validate(ks2_payload()))

    values = cell_values(load(render_ks2(act)))

    assert float(act.items[0].amount) in values
    assert float(act.items[1].amount) in values
    assert any(
        isinstance(value, str) and "с коэффициентом 1,15" in value for value in values
    )


def test_ks2_without_vat():
    payload = ks2_payload(vat_mode="none")
    act = build_act(KS2Input.model_validate(payload))

    values = cell_values(load(render_ks2(act)))

    assert "Сумма НДС (без НДС)" in values


def test_ks3_sheet_structure():
    act = build_act(KS2Input.model_validate(ks2_payload()))
    certificate = build_certificate_from_acts(
        KS3FromActs(act_ids=[act.id], document_number="1", document_date=date(2026, 3, 31)),
        [act],
    )

    sheet = load(render_ks3(certificate))
    values = cell_values(sheet)

    assert sheet.title == "КС-3"
    assert "СПРАВКА О СТОИМОСТИ ВЫПОЛНЕННЫХ РАБОТ И ЗАТРАТ" in values
    assert "0322001" in values
    assert "с начала проведения работ" in values
    assert "с начала года" in values
    assert "в том числе за отчетный период" in values
    assert "Корпус 1" in values
    assert "М.П." in values
    assert float(certificate.totals.gross) in values


def test_money_cells_use_accounting_format():
    act = build_act(KS2Input.model_validate(ks2_payload()))
    sheet = load(render_ks2(act))

    amounts = [
        cell
        for row in sheet.iter_rows()
        for cell in row
        if isinstance(cell.value, float) and cell.value > 1000
    ]

    assert amounts
    assert all(cell.number_format == "#,##0.00" for cell in amounts)
